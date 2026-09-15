"""kb_extract: the Word / PowerPoint / Excel / WebVTT readers ported from the RAG branch (K7-05).

OFFLINE. Every fixture is built in memory with the same library that reads it, so no binary lives in
the repo and nothing touches GCS. Each format must (1) be accepted where it used to be refused by
name, (2) write its locator into the text the way PDF writes `[page N]`, and (3) refuse garbage
loudly instead of storing mojibake.
"""
import io
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import kb_extract as X  # noqa: E402


def _docx_bytes():
    import docx
    d = docx.Document()
    d.add_heading("Budget rules", level=1)
    d.add_paragraph("Hold LinkedIn at eighteen thousand a month until the Q4 review.")
    d.add_heading("Reporting", level=2)
    d.add_paragraph("Weekly on Mondays.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Channel", "Budget"
    t.cell(1, 0).text, t.cell(1, 1).text = "LinkedIn", "18000"
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def _pptx_bytes():
    from pptx import Presentation
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Q3 media plan"
    s1.placeholders[1].text = "Reddit AlwaysOn26 continues"
    s1.notes_slide.notes_text_frame.text = "mention the CTR dip"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Risks"
    s2.placeholders[1].text = "Trade Desk conversions do not report upstream"
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


def _xlsx_bytes(rows=5):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Budget"
    ws.append(["Channel", "Spend"])
    for i in range(rows):
        ws.append(["LinkedIn", 1000 + i])
    ws2 = wb.create_sheet("Notes"); ws2.append(["Pacing reviewed weekly"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


VTT = b"""WEBVTT

00:00:01.000 --> 00:00:04.000
<v Ian>For ResetData the LinkedIn CTR dropped to point four.

00:00:05.000 --> 00:00:08.000
<v Priya>We agreed to hold spend at eighteen k.

00:02:10.000 --> 00:02:12.000
<v Ian>Reddit review next month.
"""

SRT = b"""1
00:00:01,000 --> 00:00:04,000
Hello from an SRT file.

2
00:00:05,000 --> 00:00:06,000
Second cue.
"""


class Formats(unittest.TestCase):
    def test_docx_headings_paragraphs_and_tables(self):
        got = X.extract("rules.docx", _docx_bytes())
        t = got["text"]
        self.assertIn("## Budget rules", t)
        self.assertIn("eighteen thousand", t)
        self.assertIn("[table 1]", t)
        self.assertIn("LinkedIn | 18000", t)
        self.assertEqual(got["kind_hint"], "reference")

    def test_pptx_one_block_per_slide_with_notes(self):
        got = X.extract("plan.pptx", _pptx_bytes())
        t = got["text"]
        self.assertIn("[slide 1]", t)
        self.assertIn("[slide 2]", t)
        self.assertIn("Q3 media plan", t)
        self.assertIn("Speaker notes: mention the CTR dip", t)
        self.assertLess(t.index("[slide 1]"), t.index("[slide 2]"))
        self.assertEqual(got["pages"], 2)
        self.assertEqual(got["kind_hint"], "plan")

    def test_xlsx_one_block_per_sheet_values_not_formulas(self):
        got = X.extract("budget.xlsx", _xlsx_bytes())
        t = got["text"]
        self.assertIn("[sheet Budget]", t)
        self.assertIn("[sheet Notes]", t)
        self.assertIn("Channel | Spend", t)
        self.assertIn("LinkedIn | 1000", t)

    def test_xlsx_long_sheet_is_cut_and_the_cut_is_declared(self):
        with unittest_patch(X, "MAX_SHEET_ROWS", 3):
            got = X.extract("big.xlsx", _xlsx_bytes(rows=10))
        self.assertIn("[Import note: sheet 'Budget' was cut to fit", got["text"])
        self.assertIn("Anything past this point is unknown, not absent", got["text"])

    def test_vtt_groups_cues_by_time_and_keeps_speakers(self):
        got = X.extract("call.vtt", VTT)
        t = got["text"]
        self.assertIn("[00:00:01]", t)
        self.assertIn("[00:02:10]", t)          # second group starts >= 120 s later
        self.assertIn("Ian: For ResetData", t)
        self.assertIn("Priya: We agreed", t)
        self.assertNotIn("<v ", t)
        self.assertEqual(got["kind_hint"], "meeting")

    def test_srt_is_read_too(self):
        t = X.extract("call.srt", SRT)["text"]
        self.assertIn("[00:00:01]", t)
        self.assertIn("Second cue.", t)

    def test_garbage_in_each_format_is_refused_not_stored(self):
        for name in ("x.docx", "x.pptx", "x.xlsx"):
            with self.assertRaises(X.ExtractError, msg=name):
                X.extract(name, b"\x00\x01\x02 not a zip container at all")
        with self.assertRaises(X.ExtractError):
            X.extract("x.vtt", b"just prose, no cues anywhere")

    def test_legacy_binaries_still_refused_by_name(self):
        for name in ("a.doc", "a.xls", "a.ppt", "a.zip"):
            with self.assertRaises(X.ExtractError) as cm:
                X.extract(name, b"xx")
            self.assertIn("not read", str(cm.exception))

    def test_title_skips_every_locator_marker(self):
        self.assertEqual(X.title_for("", "[slide 1]\nQ3 media plan"), "Q3 media plan")
        self.assertEqual(X.title_for("", "[00:00:01]\nIan: hello"), "Ian: hello")

    def test_text_and_pdf_paths_are_unchanged(self):
        got = X.extract("n.md", b"# A note\n\nbody")
        self.assertEqual(got["text"], "# A note\n\nbody")
        with self.assertRaises(X.ExtractError):
            X.extract("n.pdf", b"not a pdf")


class unittest_patch:
    """tiny attribute patcher - keeps the test file free of mock for one constant"""
    def __init__(self, obj, name, val):
        self.obj, self.name, self.val = obj, name, val

    def __enter__(self):
        self.old = getattr(self.obj, self.name); setattr(self.obj, self.name, self.val)

    def __exit__(self, *a):
        setattr(self.obj, self.name, self.old)


if __name__ == "__main__":
    unittest.main()
